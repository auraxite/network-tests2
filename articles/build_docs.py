#!/usr/bin/env python3
"""Сборка статьи и тезисов для конференции "Научный сервис в сети Интернет".

За основу берётся официальный шаблон abrau-2026.docx: он открывается,
из него удаляется всё содержимое тела, после чего наполняется своим текстом.
Такой способ сохраняет ровно тот набор стилей, поля страницы и нумерацию,
что заданы организаторами, и не добавляет ни одного нового стиля — чего
шаблон явно требует.

Имена стилей берутся из шаблона как есть. Обрати внимание: англоязычный
заголовок там называется "Згл-агнл" (именно с опечаткой), хотя в тексте
инструкции упомянут как "Згл-англ".
"""

import shutil
from datetime import datetime
from pathlib import Path

from docx import Document
from docx.shared import Cm

HERE = Path(__file__).resolve().parent
TEMPLATE = HERE / "abrau-2026.docx"
OUT_DIR = HERE / "output"

NBSP = " "  # неразрывный пробел: между инициалами и фамилией


# --------------------------------------------------------------------------
# Инфраструктура
# --------------------------------------------------------------------------

def blank_from_template() -> Document:
    """Шаблон с очищенным телом: стили, поля и нумерация сохраняются."""
    doc = Document(str(TEMPLATE))
    body = doc.element.body
    for child in list(body):
        # sectPr хранит поля страницы и ориентацию — его оставляем
        if child.tag.endswith("}sectPr"):
            continue
        body.remove(child)
    return doc


def backup(path):
    """Сохранить копию перед перезаписью.

    Документы правятся вручную в Word, а сборка затирает файл целиком.
    Поэтому каждая пересборка сначала откладывает предыдущую версию в
    output/_backup/ с отметкой времени — ручные правки можно будет достать.
    """
    if not path.exists():
        return None
    bdir = path.parent / "_backup"
    bdir.mkdir(exist_ok=True)
    stamp = datetime.fromtimestamp(path.stat().st_mtime).strftime("%Y%m%d-%H%M%S")
    dst = bdir / f"{path.stem}-{stamp}{path.suffix}"
    if not dst.exists():
        shutil.copy2(path, dst)
    return dst


def para(doc, style, text=""):
    p = doc.add_paragraph(style=style)
    if text:
        p.add_run(text)
    return p


def para_sup(doc, style, parts):
    """Абзац со смешанными обычными и надстрочными фрагментами.

    parts — список (текст, надстрочный?). Надстрочные цифры у фамилий и
    организаций шаблон требует оформлять именно шрифтовым надстрочным
    начертанием, а не символами Unicode.
    """
    p = doc.add_paragraph(style=style)
    for text, sup in parts:
        run = p.add_run(text)
        if sup:
            run.font.superscript = True
    return p


def figure(doc, image_path, caption, width_cm=14.0):
    p = doc.add_paragraph(style="Рисунок")
    p.add_run().add_picture(str(image_path), width=Cm(width_cm))
    para(doc, "Подрисуночная", caption)


# --------------------------------------------------------------------------
# Общие данные обеих работ
# --------------------------------------------------------------------------

TITLE_RU = ("Разработка методов определения задержек при нагрузочном "
            "тестировании при передаче данных между графическими процессорами "
            "в вычислительном кластере")

TITLE_EN = ("Development of methods for measuring latency during data transfer "
            "between graphics processing units in a computing cluster")

AUTHORS_RU = [
    (f"К.{NBSP}О.{NBSP}Кибизов", "1"),
    (f"А.{NBSP}Н.{NBSP}Сальников", "1,2"),
    (f"Г.{NBSP}М.{NBSP}Михайлов", "2"),
]
AUTHORS_EN = [
    (f"K.{NBSP}O.{NBSP}Kibizov", "1"),
    (f"A.{NBSP}N.{NBSP}Salnikov", "1,2"),
    (f"G.{NBSP}M.{NBSP}Mikhaylov", "2"),
]

AFFIL_RU = [
    ("1", "МГУ имени М. В. Ломоносова, Москва"),
    ("2", "Вычислительный центр им. А. А. Дородницына ФИЦ ИУ РАН, Москва"),
]
AFFIL_EN = [
    ("1", "Lomonosov Moscow State University, Moscow"),
    ("2", "Dorodnicyn Computing Centre FRC CSC RAS, Moscow"),
]

KEYWORDS_RU = ("Ключевые слова: вычислительный кластер, графический процессор "
               "(GPU), задержка передачи, нагрузочное тестирование, "
               "коммуникационная среда, привязка ресурсов, CUDA, MPI, UCX")
KEYWORDS_EN = ("Keywords: computing cluster, graphics processing unit (GPU), "
               "transfer latency, load testing, communication environment, "
               "resource binding, CUDA, MPI, UCX")


def authors_block(doc, authors):
    parts = []
    for i, (name, mark) in enumerate(authors):
        if i:
            parts.append((", ", False))
        parts.append((name, False))
        parts.append((mark, True))
    para_sup(doc, "Автор", parts)


def affil_block(doc, affils):
    for mark, name in affils:
        para_sup(doc, "Откуда", [(mark + " ", True), (name, False)])


REFS_RU = [
    "1.\tСальников А.Н., Бегаев А.А., Майсурадзе А.И. Кластеризация задержек "
    "между узлами суперкомпьютера и визуализация полученных результатов // "
    "Современные проблемы математического моделирования, обработки изображений "
    "и параллельных вычислений 2017: труды Международной научной конференции. "
    "— Ростов-на-Дону, 2017.",

    "2.\tБегаев А.А., Сальников А.Н. Метод измерения задержек при передаче "
    "данных между графическими процессорами, находящимися на разных узлах "
    "вычислительного кластера. — URL: https://cyberleninka.ru/article/n/"
    "metod-izmereniya-zaderzhek-pri-peredache-dannyh-mezhdu-graficheskimi-"
    "protsessorami-nahodyaschimisya-na-raznyh-uzlah (дата обращения: 20.09.2026).",

    "3.\tClustbench / network-tests2. — URL: "
    "https://github.com/clustbench/network-tests2 (дата обращения: 20.09.2026).",

    "4.\tNVIDIA. CUDA C++ Programming Guide: Peer Device Memory Access. — URL: "
    "https://docs.nvidia.com/cuda/cuda-c-programming-guide/ "
    "(дата обращения: 20.09.2026).",

    "5.\tNVIDIA. GPUDirect RDMA. — URL: "
    "https://docs.nvidia.com/cuda/gpudirect-rdma/ (дата обращения: 20.09.2026).",

    "6.\tOpen MPI: CUDA-Aware Support. — URL: "
    "https://www.open-mpi.org/faq/?category=runcuda (дата обращения: 20.09.2026).",

    "7.\tUnified Communication X. — URL: https://openucx.org/ "
    "(дата обращения: 20.09.2026).",

    "8.\tHoefler T., Belli R. Scientific Benchmarking of Parallel Computing "
    "Systems // SC '15. — 2015. — DOI: 10.1145/2807591.2807644",

    "9.\tВолчанинов А.П., Сальников А.Н. Исследование структуры серии измерений "
    "задержки при нагрузочном тестировании коммутационной среды вычислительного "
    "кластера // Труды RuSCDays-2023. — 2023. — С. 42.",

    "10.\tSlurm Workload Manager: Generic Resource (GRES) Scheduling. — URL: "
    "https://slurm.schedmd.com/gres.html (дата обращения: 20.09.2026).",
]

REFS_EN = [
    "1.\tSalnikov A.N., Begaev A.A., Maisuradze A.I. Klasterizatsiia zaderzhek "
    "mezhdu uzlami superkompiutera i vizualizatsiia poluchennykh rezultatov // "
    "Sovremennye problemy matematicheskogo modelirovaniia, obrabotki izobrazhenii "
    "i parallelnykh vychislenii 2017: trudy Mezhdunarodnoi nauchnoi konferentsii. "
    "— Rostov-na-Donu, 2017.",

    "2.\tBegaev A.A., Salnikov A.N. Metod izmereniia zaderzhek pri peredache "
    "dannykh mezhdu graficheskimi protsessorami, nakhodiashchimisia na raznykh "
    "uzlakh vychislitelnogo klastera. — URL: https://cyberleninka.ru/article/n/"
    "metod-izmereniya-zaderzhek-pri-peredache-dannyh-mezhdu-graficheskimi-"
    "protsessorami-nahodyaschimisya-na-raznyh-uzlah (accessed: 20.09.2026).",

    "3.\tClustbench / network-tests2. — URL: "
    "https://github.com/clustbench/network-tests2 (accessed: 20.09.2026).",

    "4.\tNVIDIA. CUDA C++ Programming Guide: Peer Device Memory Access. — URL: "
    "https://docs.nvidia.com/cuda/cuda-c-programming-guide/ (accessed: 20.09.2026).",

    "5.\tNVIDIA. GPUDirect RDMA. — URL: "
    "https://docs.nvidia.com/cuda/gpudirect-rdma/ (accessed: 20.09.2026).",

    "6.\tOpen MPI: CUDA-Aware Support. — URL: "
    "https://www.open-mpi.org/faq/?category=runcuda (accessed: 20.09.2026).",

    "7.\tUnified Communication X. — URL: https://openucx.org/ "
    "(accessed: 20.09.2026).",

    "8.\tHoefler T., Belli R. Scientific Benchmarking of Parallel Computing "
    "Systems // SC '15. — 2015. — DOI: 10.1145/2807591.2807644",

    "9.\tVolchaninov A.P., Salnikov A.N. Issledovanie struktury serii izmerenii "
    "zaderzhki pri nagruzochnom testirovanii kommutatsionnoi sredy "
    "vychislitelnogo klastera // Trudy RuSCDays-2023. — 2023. — P. 42.",

    "10.\tSlurm Workload Manager: Generic Resource (GRES) Scheduling. — URL: "
    "https://slurm.schedmd.com/gres.html (accessed: 20.09.2026).",
]


# --------------------------------------------------------------------------
# Сборка
# --------------------------------------------------------------------------

def preamble(doc, abstract_ru, abstract_en):
    """Два предваряющих блока: русский, затем английский."""
    para(doc, "Heading 1", TITLE_RU)
    authors_block(doc, AUTHORS_RU)
    affil_block(doc, AFFIL_RU)
    para(doc, "Аннотация", abstract_ru)
    para(doc, "Ключевые", KEYWORDS_RU)

    para(doc, "Згл-агнл", TITLE_EN)
    authors_block(doc, AUTHORS_EN)
    affil_block(doc, AFFIL_EN)
    para(doc, "Аннотация", abstract_en)
    para(doc, "Ключевые", KEYWORDS_EN)


def body(doc, items, fig_path=None, fig_caption=None):
    for item in items:
        if item == "FIGURE":
            if fig_path is not None and fig_path.exists():
                figure(doc, fig_path, fig_caption)
            continue
        style, text = item
        para(doc, style, text)


def bibliography(doc, refs_ru, refs_en):
    """Список литературы.

    Каждый элемент начинается с ВЕДУЩЕГО табулятора — так устроены образцы в
    шаблоне ('\\t1.\\tАдамович…'). Стиль "Литература" задаёт два табулостопа:
    правый на 397 и левый на 510, при отступе left=510, hanging=510. Первый
    табулятор прижимает номер вправо, за счёт чего "1." и "10." выравниваются
    по точке, второй ставит текст на общую левую границу. Без ведущего
    табулятора номер съезжает влево и слипается с текстом.
    """
    para(doc, "Раздел", "Литература")
    for r in refs_ru:
        para(doc, "Литература", "\t" + r)
    para(doc, "Раздел", "References")
    for r in refs_en:
        para(doc, "Литература", "\t" + r)


def build_article():
    import content_article as C
    doc = blank_from_template()
    preamble(doc, C.ABSTRACT_RU, C.ABSTRACT_EN)
    body(doc, C.BODY, OUT_DIR / "fig1_bandwidth.png", C.FIGURE_CAPTION)
    bibliography(doc, REFS_RU, REFS_EN)
    out = OUT_DIR / "Статья.docx"
    backup(out)
    doc.save(str(out))
    return out


def build_theses():
    import content_theses as C
    doc = blank_from_template()
    preamble(doc, C.ABSTRACT_RU, C.ABSTRACT_EN)
    body(doc, C.BODY)
    bibliography(doc, C.REFS_RU, C.REFS_EN)
    out = OUT_DIR / "Тезисы.docx"
    backup(out)
    doc.save(str(out))
    return out


def main():
    OUT_DIR.mkdir(parents=True, exist_ok=True)
    for path in (build_article(), build_theses()):
        d = Document(str(path))
        words = sum(len(p.text.split()) for p in d.paragraphs)
        print(f"{path.name}: абзацев {len(d.paragraphs)}, слов {words}")


if __name__ == "__main__":
    main()
